from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import CompareResult, FileHit, FileRef, MatchRow

HEADER_FILL = PatternFill("solid", fgColor="548235")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ORANGE_FONT = Font(color="C45C26", bold=True)
THIN = Border(
    left=Side(style="thin", color="B8C4B0"),
    right=Side(style="thin", color="B8C4B0"),
    top=Side(style="thin", color="B8C4B0"),
    bottom=Side(style="thin", color="B8C4B0"),
)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _nombres(refs: list[FileRef]) -> str:
    return " | ".join(item.name for item in refs)


def _cell_files(cell, refs: list[FileRef]) -> None:
    cell.value = _nombres(refs)
    cell.alignment = LEFT
    cell.border = THIN
    if any(item.subcarpeta for item in refs):
        cell.font = ORANGE_FONT


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(1, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    ws.freeze_panes = "A2"


def _autosize(ws: Worksheet, min_width: int = 12, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_width
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(value) + 2))
        ws.column_dimensions[letter].width = width


def _sheet_rows(ws: Worksheet, rows: list[MatchRow], con_bom: bool, con_wi: bool) -> None:
    headers = ["Referencia"]
    if con_bom:
        headers.append("BOM")
    if con_wi:
        headers.append("WI")
    _write_header(ws, headers)
    for i, row in enumerate(rows, start=2):
        values: list = [row.referencia]
        if con_bom:
            values.append(row.boms)
        if con_wi:
            values.append(row.wis)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(i, col)
            if isinstance(value, list):
                _cell_files(cell, value)
            else:
                cell.value = value
                cell.alignment = LEFT
                cell.border = THIN
    last = max(2, len(rows) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    _autosize(ws)


def _sheet_files(ws: Worksheet, hits: list[FileHit], extra: str) -> None:
    _write_header(ws, ["Archivo", extra, "Ruta"])
    for i, hit in enumerate(hits, start=2):
        for col, value in enumerate((hit.name, hit.limpio, hit.path), start=1):
            cell = ws.cell(i, col, value)
            cell.alignment = LEFT
            cell.border = THIN
    last = max(2, len(hits) + 1)
    ws.auto_filter.ref = f"A1:C{last}"
    _autosize(ws, max_width=70)


def write_excel(result: CompareResult, path: str | Path) -> Path:
    out = Path(path)
    wb = Workbook()

    ws = wb.active
    ws.title = "Faltan WI"
    _sheet_rows(ws, result.faltan, con_bom=True, con_wi=False)

    ws = wb.create_sheet("Completas")
    _sheet_rows(ws, result.coinciden, con_bom=True, con_wi=True)

    ws = wb.create_sheet("WI sin BOM")
    _sheet_rows(ws, result.sobran, con_bom=False, con_wi=True)

    ws = wb.create_sheet("BOM no reconocidos")
    _sheet_files(ws, result.bom_no_reconocidos, "Limpio")

    ws = wb.create_sheet("WI no reconocidas")
    _sheet_files(ws, result.wi_no_reconocidas, "Limpio")

    if result.omitidos_bom:
        ws = wb.create_sheet("BOM omitidos")
        _sheet_files(ws, result.omitidos_bom, "Motivo")

    ws = wb.create_sheet("Resumen")
    _write_header(ws, ["Campo", "Valor"])
    resumen = [
        ("Carpeta BOM", result.carpeta_bom),
        ("Carpeta WI", result.carpeta_wi),
        ("Archivos BOM", result.n_bom_archivos),
        ("Archivos WI", result.n_wi_archivos),
        ("Referencias BOM", result.n_bom_refs),
        ("Referencias WI", result.n_wi_refs),
        ("Faltan WI", len(result.faltan)),
        ("Completas", len(result.coinciden)),
        ("WI sin BOM", len(result.sobran)),
        ("BOM no reconocidos", len(result.bom_no_reconocidos)),
        ("WI no reconocidas", len(result.wi_no_reconocidas)),
        ("BOM antiguas (subcarpeta)", result.n_bom_obsoletas),
    ]
    for i, (campo, valor) in enumerate(resumen, start=2):
        for col, value in enumerate((campo, valor), start=1):
            cell = ws.cell(i, col, value)
            cell.alignment = LEFT
            cell.border = THIN
    _autosize(ws, max_width=80)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
