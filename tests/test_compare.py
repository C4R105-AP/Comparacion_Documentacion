from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from wi_compare.compare import comparar_carpetas, comparar_nombres
from wi_compare.names import extraer_desde_bom, extraer_desde_wi, limpiar_referencia
from wi_compare.report import write_excel


class LimpiezaNombresTests(unittest.TestCase):
    def test_prefijos_bom_cad_ce(self):
        self.assertEqual(limpiar_referencia("BOM CAD CE-REF"), "REF")
        self.assertEqual(limpiar_referencia("BOM CE-BOARD-AX-001"), "BOARD-AX-001")
        self.assertEqual(limpiar_referencia("PCB-BOARD-AX-001"), "BOARD-AX-001")
        self.assertEqual(limpiar_referencia("ASSY MID-AP-130"), "MID-AP-130")

    def test_fechas_finales(self):
        self.assertEqual(limpiar_referencia("PROD-A_2020-01-29"), "PROD-A")
        self.assertEqual(limpiar_referencia("PROD-A 29-01-2020"), "PROD-A")
        self.assertEqual(limpiar_referencia("PROD-A_20140527"), "PROD-A")

    def test_fecha_y_sufijo_descriptivo(self):
        bom = extraer_desde_bom("CE-PLACA+001-AB_251109_Acabado_Especial")
        wi = extraer_desde_wi("-PLACA+001-AB - Instrucción de fabricación")
        self.assertEqual(bom, "PLACA+001-AB")
        self.assertEqual(wi, "PLACA+001-AB")
        self.assertEqual(
            extraer_desde_bom("CE-PLACA+002-025_1284M-AC_260622"),
            "PLACA+002-025-1284M-AC",
        )

    def test_separadores(self):
        self.assertEqual(limpiar_referencia("MID_AP-130 AB"), "MID-AP-130-AB")

    def test_instruccion_de_fabricacion(self):
        stem = "CE-BOARD-AX-001 Instrucción de fabricación Rev.3"
        self.assertEqual(limpiar_referencia(stem), "BOARD-AX-001")
        self.assertEqual(extraer_desde_wi(stem), "BOARD-AX-001")

    def test_wi_corta_con_codigo(self):
        self.assertEqual(extraer_desde_wi("CE-BOARD-AX-001"), "BOARD-AX-001")
        self.assertEqual(extraer_desde_wi("WI-MID-AP-130"), "MID-AP-130")

    def test_wi_sin_codigo_se_ignora(self):
        self.assertIsNone(extraer_desde_wi("Notas de montaje generales"))

    def test_coordenadas_omitidas(self):
        self.assertIsNone(extraer_desde_bom("BOM CE-REF coordenadas"))

    def test_bom_y_wi_misma_clave(self):
        bom = extraer_desde_bom("BOM CE-BOARD-AX-001_2024-03-01")
        wi = extraer_desde_wi("CE-BOARD-AX-001 - Instrucción de fabricación")
        self.assertEqual(bom, wi)
        self.assertEqual(bom, "BOARD-AX-001")


class CruceCarpetasTests(unittest.TestCase):
    def test_falta_coincide_y_sobra(self):
        with TemporaryDirectory() as tmp:
            boms = Path(tmp) / "boms"
            wis = Path(tmp) / "wis"
            boms.mkdir()
            wis.mkdir()
            (boms / "BOM CE-BOARD-AX-001.xlsx").write_bytes(b"x")
            (boms / "BOM CE-MISSING.xls").write_bytes(b"x")
            (boms / "BOM CE-REF coordenadas.xlsx").write_bytes(b"x")
            (wis / "CE-BOARD-AX-001 Instrucción de fabricación.docx").write_bytes(b"x")
            (wis / "CE-ORPHAN Instrucción de fabricación.docx").write_bytes(b"x")

            result = comparar_carpetas(str(boms), str(wis))

            self.assertEqual([r.referencia for r in result.faltan], ["MISSING"])
            self.assertEqual([r.referencia for r in result.coinciden], ["BOARD-AX-001"])
            self.assertEqual([r.referencia for r in result.sobran], ["ORPHAN"])
            self.assertEqual(len(result.omitidos_bom), 1)
            self.assertEqual(result.n_bom_archivos, 2)
            self.assertEqual(result.n_wi_archivos, 2)

            xlsx = Path(tmp) / "informe.xlsx"
            write_excel(result, xlsx)
            wb = load_workbook(xlsx)
            self.assertIn("Faltan WI", wb.sheetnames)
            self.assertEqual(wb["Faltan WI"]["A2"].value, "MISSING")
            self.assertEqual(wb["Completas"]["A2"].value, "BOARD-AX-001")


class SubcarpetasTests(unittest.TestCase):
    def test_bom_antigua_se_descarta_y_wi_en_subcarpeta_vale(self):
        with TemporaryDirectory() as tmp:
            boms = Path(tmp) / "boms"
            wis = Path(tmp) / "wis"
            (boms / "old").mkdir(parents=True)
            (wis / "old").mkdir(parents=True)
            (boms / "BOM CE-BOARD-AX-001.xlsx").write_bytes(b"x")
            (boms / "old" / "BOM CE-BOARD-AX-001.xlsx").write_bytes(b"x")
            (boms / "old" / "BOM CE-LEGACY.xls").write_bytes(b"x")
            (wis / "old" / "CE-BOARD-AX-001 Instrucción de fabricación.docx").write_bytes(b"x")

            result = comparar_carpetas(str(boms), str(wis))

            self.assertEqual([r.referencia for r in result.faltan], [])
            self.assertEqual([r.referencia for r in result.coinciden], ["BOARD-AX-001"])
            self.assertTrue(result.coinciden[0].wis[0].subcarpeta)
            self.assertFalse(result.coinciden[0].boms[0].subcarpeta)
            self.assertEqual(result.n_bom_obsoletas, 2)
            self.assertNotIn("LEGACY", [r.referencia for r in result.faltan + result.coinciden])


class CruceNombresTests(unittest.TestCase):
    def test_desde_nombres_de_carpeta(self):
        result = comparar_nombres(
            [
                {"name": "BOM CE-BOARD-AX-001.xlsx", "relative": "boms/BOM CE-BOARD-AX-001.xlsx"},
                {"name": "BOM CE-MISSING.xls", "relative": "boms/BOM CE-MISSING.xls"},
                {"name": "BOM CE-REF coordenadas.xlsx", "relative": "boms/BOM CE-REF coordenadas.xlsx"},
                {"name": "nested.xlsx", "relative": "boms/old/nested.xlsx"},
            ],
            [
                {"name": "CE-BOARD-AX-001 Instrucción de fabricación.docx", "relative": "wis/CE-BOARD-AX-001 Instrucción de fabricación.docx"},
                {"name": "CE-ORPHAN Instrucción de fabricación.docx", "relative": "wis/CE-ORPHAN Instrucción de fabricación.docx"},
            ],
        )
        self.assertEqual([r.referencia for r in result.faltan], ["MISSING"])
        self.assertEqual([r.referencia for r in result.coinciden], ["BOARD-AX-001"])
        self.assertEqual([r.referencia for r in result.sobran], ["ORPHAN"])
        self.assertEqual(len(result.omitidos_bom), 2)
        self.assertTrue(all(r.referencia != "nested" for r in result.faltan + result.coinciden))


if __name__ == "__main__":
    unittest.main()
